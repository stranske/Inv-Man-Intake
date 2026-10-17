"""Orchestration helper for v1 acceptance smoke tests."""

from __future__ import annotations

import io
import json
import math
import os
import re
import sqlite3
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError

from inv_man_intake.data.repository import CoreRepository
from inv_man_intake.extraction.confidence import (
    ThresholdConfig,
    attach_threshold_summary,
    evaluate_thresholds,
    load_threshold_config,
    select_threshold_profile,
)
from inv_man_intake.extraction.doc_type import classify_doc_type
from inv_man_intake.extraction.orchestrator import ExtractionOrchestrator
from inv_man_intake.extraction.providers.base import ExtractedDocumentResult
from inv_man_intake.extraction.service import (
    build_legacy_pdf_fixture_service,
    build_pyodide_light_service,
    extraction_service_extractor,
)
from inv_man_intake.intake.integration import (
    filesystem_content_resolver,
    register_intake_bundle_file,
)
from inv_man_intake.intake.models import IngestRecord
from inv_man_intake.intake.service import IngestionService
from inv_man_intake.intake.standard_elements import load_standard_element_library
from inv_man_intake.observability import (
    FleetRunContext,
    InMemoryTraceSink,
    TraceContext,
    TraceEvent,
    Tracer,
    build_fleet_records,
    build_summary_from_pipeline,
    child_trace_context,
    derive_trace_url,
    ensure_correlation_id,
    extract_trace_context,
    inject_trace_context,
    new_trace_context,
)
from inv_man_intake.observability.langsmith_fleet import DEFAULT_PROJECT
from inv_man_intake.performance.characterize import characterize_series, gate_scoring_submission
from inv_man_intake.performance.conflict_resolver import resolve_source_conflicts
from inv_man_intake.performance.contracts import (
    PerformancePayload,
    PerformancePoint,
    PerformanceSeries,
)
from inv_man_intake.performance.ingest import load_xlsx_timeseries
from inv_man_intake.performance.metrics import compute_metrics
from inv_man_intake.performance.normalize import normalize_payload
from inv_man_intake.queue.assignment import create_analyst_first_assignment
from inv_man_intake.scoring.contracts import ScoreComponent, ScoreSubmission
from inv_man_intake.scoring.engine import compute_score
from inv_man_intake.scoring.explainability import (
    ComponentProvenance,
    ScoreComponentInput,
    build_explainability_payload,
    format_explainability_payload,
)
from inv_man_intake.scoring.weights import (
    get_weight_set,
    normalize_asset_class,
    weights_for_registry,
)
from inv_man_intake.storage.document_store import InMemoryDocumentStore

# Single source of truth for the default extraction threshold policy. The production CLI
# (run.py) loads this same file; the smoke/demo fallback must not drift from it (see #694).
DEFAULT_THRESHOLD_CONFIG_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "extraction_thresholds.yaml"
)
DEFAULT_STANDARD_ELEMENT_LIBRARY_PATH = (
    Path(__file__).resolve().parents[2] / "config" / "standard_elements" / "_stub.json"
)


class _FanoutTraceSink:
    """Emit trace events to all sinks while keeping tracing non-blocking."""

    def __init__(self, *sinks: object) -> None:
        self._sinks = sinks

    def on_span_start(self, event: TraceEvent) -> None:
        for sink in self._sinks:
            on_span_start = getattr(sink, "on_span_start", None)
            if callable(on_span_start):
                try:
                    on_span_start(event)
                except Exception:
                    continue

    def on_span_end(self, event: TraceEvent) -> None:
        for sink in self._sinks:
            on_span_end = getattr(sink, "on_span_end", None)
            if callable(on_span_end):
                try:
                    on_span_end(event)
                except Exception:
                    continue


@dataclass(frozen=True)
class V1SmokeArtifacts:
    service: IngestionService
    core_repository: CoreRepository
    document_store: InMemoryDocumentStore
    registration: object
    record: object
    sink: InMemoryTraceSink
    trace_context: object
    correlation_id: str
    intake_start: object
    extraction_start: object
    secondary_extraction_result: object
    performance_start: object
    threshold_decision: object
    extraction_with_thresholds: ExtractedDocumentResult
    conflict_result: object
    normalized: object
    metrics: object
    queue_assignment: object
    score: object
    formatted_explainability: dict[str, object]
    langsmith_fleet_records: list[dict[str, Any]]
    performance: dict[str, Any]


def run_v1_smoke_pipeline(
    *,
    fixture_root: Path,
    intake_bundle_file: str = "pdf_primary_mixed_bundle.json",
    package_id: str,
    expected_document_ids: tuple[str, ...],
) -> V1SmokeArtifacts:
    """Run the smoke pipeline with strict package/document-id assertions.

    Thin wrapper that delegates to :func:`_run_pipeline_core`, the single
    orchestration code path also used by the headless ``run_pipeline`` entry
    point in :mod:`inv_man_intake.run`. Keeping one core means the operator
    CLI and the acceptance smoke exercise identical pipeline behavior.
    """

    return _run_pipeline_core(
        fixture_root=fixture_root,
        intake_bundle_file=intake_bundle_file,
        package_id=package_id,
        expected_document_ids=expected_document_ids,
    )


def _run_pipeline_core(
    *,
    fixture_root: Path,
    intake_bundle_file: str = "pdf_primary_mixed_bundle.json",
    package_id: str,
    expected_document_ids: tuple[str, ...] | None = None,
    threshold_config: ThresholdConfig | None = None,
    content_root: Path | None = None,
) -> V1SmokeArtifacts:
    """Execute the deterministic intake-to-scoring pipeline once.

    When ``expected_document_ids`` is provided (the smoke path), the registered
    package state is asserted exactly. When it is ``None`` (the headless
    ``run_pipeline`` path), registration is validated softly: a rejected bundle
    raises :class:`ValueError` instead of asserting, so the CLI can surface a
    clean non-zero exit rather than an ``AssertionError`` traceback.
    """

    smoke_mode = expected_document_ids is not None
    bundle_path = fixture_root / intake_bundle_file
    parsed_bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    raw_files = parsed_bundle.get("files")
    file_entries = raw_files if isinstance(raw_files, list) else []
    resolved_content_root = content_root or fixture_root
    content_resolver = None if smoke_mode else filesystem_content_resolver(resolved_content_root)

    sink = InMemoryTraceSink()
    tracer = _entrypoint_tracer(sink=sink)
    correlation_id = ensure_correlation_id()
    trace_tags = {
        "package_id": package_id,
        "stage": "intake",
        "correlation_id": correlation_id,
    }
    trace_tags.update(_langsmith_context_tags())
    trace_context = new_trace_context(tags=trace_tags)

    service = IngestionService()
    core_repository = CoreRepository(sqlite3.connect(":memory:"))
    document_store = InMemoryDocumentStore()
    with tracer.start_span(
        name="v1_acceptance.intake_register",
        context=trace_context,
        metadata={"fixture": intake_bundle_file},
    ):
        registration = register_intake_bundle_file(
            bundle_path,
            service,
            core_repository=core_repository,
            document_store=document_store,
            content_resolver=content_resolver,
        )

    if expected_document_ids is None:
        if not registration.accepted:
            rejected = ", ".join(issue.code for issue in registration.errors) or "unknown"
            raise ValueError(f"intake bundle registration rejected: {rejected}")
        record = service.get_record(package_id)
    else:
        record = _assert_registered_package_state(
            service=service,
            package_id=package_id,
            expected_document_ids=expected_document_ids,
        )
    intake_start = _start_event(sink, "v1_acceptance.intake_register")
    extracted_context = extract_trace_context(inject_trace_context(trace_context))
    assert extracted_context is not None
    extraction_context = child_trace_context(
        extracted_context,
        parent_span_id=intake_start.span_id,
        tags={"stage": "extract"},
    )

    primary_document = core_repository.get_document(record.document_ids[0])
    assert primary_document is not None, "primary document must be registered"
    primary_file_name = primary_document.file_name
    primary_content = _pipeline_document_bytes(
        document_id=record.document_ids[0],
        fund_id=record.fund_id,
        file_name=primary_file_name,
        smoke_mode=smoke_mode,
        fixture_root=fixture_root,
        document_store=document_store,
    )
    extraction_result = _run_extraction_smoke(
        tracer=tracer,
        trace_context=extraction_context,
        source_doc_id=record.document_ids[0],
        primary_file_name=primary_file_name,
        content=primary_content,
        correlation_id=correlation_id,
    )
    if len(record.document_ids) < 2:
        raise ValueError("intake bundle must register at least two documents for extraction")
    secondary_document = core_repository.get_document(record.document_ids[1])
    assert secondary_document is not None, "secondary document must be registered"
    if smoke_mode:
        # Smoke bundles may name synthetic secondary files that are not checked in;
        # the boundary exercise only needs representative xlsx bytes.
        secondary_content = _fixture_bytes(
            fixture_root=fixture_root,
            file_name="summit_arc_track_record.xlsx",
        )
    else:
        secondary_content = _pipeline_document_bytes(
            document_id=record.document_ids[1],
            fund_id=record.fund_id,
            file_name=secondary_document.file_name,
            smoke_mode=smoke_mode,
            fixture_root=fixture_root,
            document_store=document_store,
        )
    secondary_extraction_result = _run_secondary_extraction_boundary_smoke(
        tracer=tracer,
        trace_context=extraction_context,
        source_doc_id=record.document_ids[1],
        content=secondary_content,
        correlation_id=correlation_id,
    )
    with tracer.start_span(
        name="v1_acceptance.threshold_handling",
        context=extraction_context,
        metadata={"package_id": record.package_id},
    ):
        effective_threshold_config = threshold_config or load_threshold_config(
            DEFAULT_THRESHOLD_CONFIG_PATH
        )
        threshold_key_fields, threshold_profile_config = select_threshold_profile(
            document_type=classify_doc_type(primary_content),
            key_fields=(
                "strategy.asset_class",
                "terms.management_fee",
                "performance.net_return_1y",
                "operations.aum",
                "team.key_person_risk",
            ),
            config=effective_threshold_config,
        )
        threshold_decision = evaluate_thresholds(
            result=extraction_result,
            key_fields=threshold_key_fields,
            config=threshold_profile_config,
        )
        extraction_with_thresholds = attach_threshold_summary(
            result=extraction_result,
            decision=threshold_decision,
        )

    extraction_start = _start_event(sink, "extraction_orchestrator.run")
    performance_context = child_trace_context(
        extraction_context,
        parent_span_id=extraction_start.span_id,
        tags={"stage": "performance"},
    )
    with tracer.start_span(
        name="v1_acceptance.performance_normalize",
        context=performance_context,
        metadata={"package_id": record.package_id},
    ):
        xlsx_series, deck_series, benchmark_series = _resolve_performance_series(
            smoke_mode=smoke_mode,
            file_entries=file_entries,
            record=record,
            fund_id=record.fund_id,
            content_root=resolved_content_root,
            document_store=document_store,
            core_repository=core_repository,
        )
        conflict_result = normalized = metrics = characterization = None
        if xlsx_series is not None:
            conflict_result = resolve_source_conflicts(
                xlsx_series=xlsx_series,
                other_series=deck_series,
            )
            normalized = normalize_payload(
                PerformancePayload(monthly=conflict_result.resolved_series)
            )
            metrics = compute_metrics(
                PerformancePayload(monthly=normalized.monthly),
                benchmark_monthly=benchmark_series,
            )
            characterization = characterize_series(
                normalized.monthly,
                metrics,
                source_names=_document_source_names(
                    repository=core_repository,
                    document_ids=record.document_ids,
                ),
                standard_library=load_standard_element_library(
                    DEFAULT_STANDARD_ELEMENT_LIBRARY_PATH
                ),
            )
        performance = {
            "status": "available" if xlsx_series is not None else "unavailable",
            "source": "smoke_fixture" if smoke_mode else "submitted_workbook",
            "monthly": (
                [
                    {"as_of": point.as_of.isoformat(), "value": point.value}
                    for point in normalized.monthly.points
                ]
                if normalized is not None
                else None
            ),
            "metrics": metrics.to_canonical_dict() if metrics is not None else None,
            "comparison_available": deck_series is not None,
            "benchmark_available": benchmark_series is not None,
        }

    performance_start = _start_event(sink, "v1_acceptance.performance_normalize")
    queue_context = child_trace_context(
        performance_context,
        parent_span_id=performance_start.span_id,
        tags={"stage": "queue"},
    )
    with tracer.start_span(
        name="v1_acceptance.queue_audit_output",
        context=queue_context,
        metadata={"package_id": record.package_id},
    ):
        queue_assignment = create_analyst_first_assignment(
            item_id=f"{record.package_id}:validation:performance_conflict:{correlation_id}",
            analyst_id="analyst_001",
            created_at=datetime(2026, 3, 4, 10, 0, tzinfo=UTC),
        )

    scoring_context = child_trace_context(
        performance_context,
        parent_span_id=performance_start.span_id,
        tags={"stage": "score"},
    )
    with tracer.start_span(
        name="v1_acceptance.scoring_compute",
        context=scoring_context,
        metadata={
            "manager_id": record.fund_id,
            "queue_item_id": queue_assignment.item_id,
            "correlation_id": correlation_id,
        },
    ):
        # Reference component values are smoke-only. Production scoring derives
        # components from submitted evidence (see
        # ``_evidence_backed_score_components``) and stays unavailable when that
        # evidence is below the documented minimum thresholds.
        score = None
        formatted_explainability = _unavailable_explainability(
            "evidence_backed_score_components_unavailable"
        )
        if not smoke_mode:
            derived = _evidence_backed_score_components(
                extraction=extraction_with_thresholds,
                metrics=metrics,
                workbook_ref=_workbook_provenance_ref(
                    repository=core_repository,
                    record=record,
                    file_entries=file_entries,
                ),
            )
            if derived is not None and characterization is not None:
                try:
                    production_submission = gate_scoring_submission(
                        ScoreSubmission(
                            manager_id=record.fund_id,
                            asset_class=derived.asset_class,
                            components=derived.components,
                        ),
                        characterization=characterization,
                    )
                except PermissionError:
                    production_submission = None
                    formatted_explainability = _unavailable_explainability(
                        "operator_confirmation_required"
                    )
                if production_submission is not None:
                    score = compute_score(
                        production_submission,
                        weights_by_asset_class=weights_for_registry(),
                    )
                    production_explainability = build_explainability_payload(
                        components=_explainability_inputs(
                            score.asset_class,
                            derived.components,
                            rationales=derived.rationales,
                            provenances=derived.provenances,
                        ),
                        overall_score=score.final_score,
                    )
                    formatted_explainability = format_explainability_payload(
                        production_explainability
                    )
        if smoke_mode:
            assert metrics is not None and characterization is not None
            components = _score_components(metrics.benchmark_correlation)
            submission = gate_scoring_submission(
                ScoreSubmission(
                    manager_id=record.fund_id,
                    asset_class="credit",
                    components=components,
                ),
                characterization=characterization,
            )
            score = compute_score(
                submission,
                weights_by_asset_class=weights_for_registry(),
            )
            explainability = build_explainability_payload(
                components=_explainability_inputs(score.asset_class, components),
                overall_score=score.final_score,
            )
            formatted_explainability = format_explainability_payload(explainability)
    fleet_summary = build_summary_from_pipeline(
        document_ids=record.document_ids,
        extraction=extraction_with_thresholds,
        secondary_extraction=secondary_extraction_result,
        validation_status="escalated" if threshold_decision.escalate else "accepted",
        score_count=len(score.contributions) if score is not None else 0,
        review_queue_outcome=queue_assignment.owner_role,
        artifact_refs=(
            f"artifact:packages/{record.package_id}/metadata.json",
            "artifact:extraction/threshold-summary.json",
            "artifact:scoring/explainability.json",
        ),
        trace_refs=(f"trace:{trace_context.trace_id}",),
        document_types=_derive_document_types(
            repository=core_repository, document_ids=record.document_ids
        ),
    )
    error_category = (
        threshold_decision.escalation_reason or "unknown_escalation_reason"
        if threshold_decision.escalate
        else "none"
    )

    langsmith_fleet_records = build_fleet_records(
        context=FleetRunContext(
            run_id=trace_context.run_id or trace_context.trace_id,
            package_id=record.package_id,
            provider=extraction_with_thresholds.provider_name,
            model="deterministic-pdf-parser",
            trace_id=trace_context.trace_id,
            trace_url=derive_trace_url(trace_context.trace_id),
            correlation_id=correlation_id,
            latency_ms=_pipeline_latency_ms(
                sink=sink, root_span_name="v1_acceptance.intake_register"
            ),
            error_category=error_category,
        ),
        summary=fleet_summary,
        artifact_ref="artifact:langsmith-fleet.ndjson",
    )

    return V1SmokeArtifacts(
        service=service,
        core_repository=core_repository,
        document_store=document_store,
        registration=registration,
        record=record,
        sink=sink,
        trace_context=trace_context,
        correlation_id=correlation_id,
        intake_start=intake_start,
        extraction_start=extraction_start,
        secondary_extraction_result=secondary_extraction_result,
        performance_start=performance_start,
        threshold_decision=threshold_decision,
        extraction_with_thresholds=extraction_with_thresholds,
        conflict_result=conflict_result,
        normalized=normalized,
        metrics=metrics,
        queue_assignment=queue_assignment,
        score=score,
        formatted_explainability=formatted_explainability,
        langsmith_fleet_records=langsmith_fleet_records,
        performance=performance,
    )


def _entrypoint_tracer(*, sink: InMemoryTraceSink) -> Tracer:
    tracer = Tracer(enabled=True, sink=sink)
    if not os.getenv("LANGSMITH_API_KEY", "").strip():
        return tracer
    from inv_man_intake.observability.langsmith_fleet import ensure_langsmith_project_defaults
    from inv_man_intake.observability.langsmith_sink import LangSmithTraceSink

    ensure_langsmith_project_defaults()
    return Tracer(
        enabled=True,
        sink=_FanoutTraceSink(
            sink,
            LangSmithTraceSink.from_env(),
        ),
    )


def _langsmith_context_tags() -> dict[str, str]:
    if not os.getenv("LANGSMITH_API_KEY", "").strip():
        return {}
    project_name = os.getenv("LANGSMITH_PROJECT", "").strip() or DEFAULT_PROJECT
    return {
        "langsmith_enabled": "true",
        "langsmith_project": project_name,
    }


def _run_extraction_smoke(
    *,
    tracer: Tracer,
    trace_context: TraceContext,
    source_doc_id: str,
    primary_file_name: str,
    content: bytes,
    correlation_id: str,
) -> ExtractedDocumentResult:
    # This historical smoke corpus contains a deliberately minimal pseudo-PDF
    # understood by pdf-primary, not a conforming PDF for the shared extractor.
    # Keep its golden baseline explicit; the default service selects Doc-Lineage.
    service = (
        build_legacy_pdf_fixture_service()
        if primary_file_name.lower().endswith(".pdf")
        else build_pyodide_light_service(primary_file_name)
    )
    orchestrator = ExtractionOrchestrator(
        primary_name=service.backend_name,
        fallback_name="fixture-fallback",
        fallback_extractor=lambda payload: {"document_id": payload["document_id"]},
        primary_extractor=extraction_service_extractor(service),
        tracer=tracer,
    )
    result = orchestrator.run(
        {
            "id": f"{source_doc_id}:extract",
            "document_id": source_doc_id,
            "content": content,
            "correlation_id": correlation_id,
        },
        trace_context=trace_context,
    )
    assert result.resolved is True
    assert result.provider_used == service.backend_name
    assert result.data is not None
    extracted = result.data["result"]
    assert isinstance(extracted, ExtractedDocumentResult)
    return extracted


def _run_secondary_extraction_boundary_smoke(
    *,
    tracer: Tracer,
    trace_context: TraceContext,
    source_doc_id: str,
    content: bytes,
    correlation_id: str,
) -> object:
    service = build_pyodide_light_service("boundary.pdf")
    extractor_key = "_".join(("primary", "extractor"))
    orchestrator = ExtractionOrchestrator(
        primary_name=service.backend_name,
        fallback_name="secondary-unsupported-escalation",
        fallback_extractor=_unsupported_secondary_extractor,
        tracer=tracer,
        **{extractor_key: extraction_service_extractor(service)},  # type: ignore[arg-type]
    )
    result = orchestrator.run(
        {
            "id": f"{source_doc_id}:secondary-extract",
            "document_id": source_doc_id,
            "content": content,
            "correlation_id": correlation_id,
        },
        trace_context=trace_context,
    )
    assert result.resolved is False
    assert result.escalation_route == "ops_review"
    return result


def _unsupported_secondary_extractor(payload: dict[str, Any]) -> dict[str, object]:
    content = payload["content"]
    if not isinstance(content, (bytes, bytearray)):
        raise TypeError("document content must be bytes")
    raise ValueError(_unsupported_secondary_bytes_reason(bytes(content)))


def _fixture_bytes(*, fixture_root: Path, file_name: str) -> bytes:
    content_root = (fixture_root.parent / "extraction").resolve()
    source_path = (content_root / file_name).resolve()
    try:
        source_path.relative_to(content_root)
    except ValueError as exc:
        raise ValueError(
            f"intake entry file_name {file_name!r} escapes the content base directory"
        ) from exc
    return source_path.read_bytes()


def _pipeline_document_bytes(
    *,
    document_id: str,
    fund_id: str,
    file_name: str,
    smoke_mode: bool,
    fixture_root: Path,
    document_store: InMemoryDocumentStore,
) -> bytes:
    if smoke_mode:
        return _fixture_bytes(fixture_root=fixture_root, file_name=file_name)
    document_key = f"{fund_id}/{document_id}"
    versions = document_store.list_versions(document_key)
    if not versions:
        raise ValueError(f"no persisted document bytes for {document_id}")
    return document_store.get(document_key, versions[-1].version_id)


def _bundle_has_performance_track_record(file_entries: list[dict[str, Any]]) -> bool:
    return any(
        isinstance(entry, dict) and entry.get("role") == "performance_track_record"
        for entry in file_entries
    )


def _performance_track_record_file_name(file_entries: list[dict[str, Any]]) -> str | None:
    for entry in file_entries:
        if isinstance(entry, dict) and entry.get("role") == "performance_track_record":
            file_name = entry.get("file_name")
            if isinstance(file_name, str) and file_name.strip():
                return file_name
    return None


def _document_id_for_file_name(
    *,
    repository: CoreRepository,
    document_ids: tuple[str, ...] | list[str],
    file_name: str,
) -> str | None:
    for document_id in document_ids:
        document = repository.get_document(document_id)
        if document is not None and document.file_name == file_name:
            return document_id
    return None


def _require_valid_performance_track_record_bytes(content: bytes, *, file_name: str) -> None:
    if not content.startswith(b"PK\x03\x04"):
        raise ValueError(
            f"performance track record {file_name} is not a valid xlsx container: "
            f"{_unsupported_secondary_bytes_reason(content)}"
        )
    kind = _ooxml_zip_kind(content)
    if kind != "xlsx":
        raise ValueError(f"performance track record {file_name} must be xlsx bytes, got {kind}")


def _resolve_performance_series(
    *,
    smoke_mode: bool,
    file_entries: list[dict[str, Any]],
    record: IngestRecord | None = None,
    fund_id: str | None = None,
    content_root: Path | None = None,
    document_store: InMemoryDocumentStore | None = None,
    core_repository: CoreRepository | None = None,
) -> tuple[PerformanceSeries | None, PerformanceSeries | None, PerformanceSeries | None]:
    if smoke_mode:
        return _fixture_performance_series()
    if _bundle_has_performance_track_record(file_entries):
        track_record_name = _performance_track_record_file_name(file_entries)
        if track_record_name is None:
            raise ValueError(
                "performance data unavailable: bundle declares performance_track_record without file_name"
            )
        if (
            record is None
            or fund_id is None
            or content_root is None
            or document_store is None
            or core_repository is None
        ):
            raise ValueError(
                "performance data unavailable: headless performance resolution requires pipeline context"
            )
        document_id = _document_id_for_file_name(
            repository=core_repository,
            document_ids=record.document_ids,
            file_name=track_record_name,
        )
        if document_id is None:
            raise ValueError(
                f"performance data unavailable: no registered document for {track_record_name}"
            )
        track_record_bytes = _pipeline_document_bytes(
            document_id=document_id,
            fund_id=fund_id,
            file_name=track_record_name,
            smoke_mode=False,
            fixture_root=content_root,
            document_store=document_store,
        )
        _require_valid_performance_track_record_bytes(
            track_record_bytes,
            file_name=track_record_name,
        )
        return _workbook_performance_series(track_record_bytes), None, None
    return None, None, None


def _workbook_performance_series(content: bytes) -> PerformanceSeries | None:
    """Read canonical XLSX rows without guessing units or selecting ambiguous tables.

    One sheet must have `as_of` and `value` headers, with optional `frequency`
    (defaults to monthly). Values are decimal returns; dates are ISO text or
    Excel dates. Unsupported, malformed, or ambiguous workbooks are unavailable.
    """
    from openpyxl import load_workbook  # type: ignore[import-untyped]
    from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            candidates = []
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                names = [str(cell).strip().lower() if cell is not None else "" for cell in header]
                if not {"as_of", "value"}.issubset(names):
                    continue
                if any(names.count(key) > 1 for key in ("as_of", "value", "frequency")):
                    return None
                parsed = []
                for values in rows:
                    if all(value is None for value in values):
                        continue
                    row = dict(zip(names, values, strict=True))
                    as_of = row.get("as_of")
                    if isinstance(as_of, datetime):
                        as_of = as_of.date()
                    if isinstance(as_of, date):
                        as_of = as_of.isoformat()
                    parsed.append(
                        {
                            "frequency": row.get("frequency", "monthly"),
                            "as_of": as_of,
                            "value": row.get("value"),
                        }
                    )
                candidates.append(load_xlsx_timeseries(parsed).monthly)
            return candidates[0] if len(candidates) == 1 else None
        finally:
            workbook.close()
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        zipfile.BadZipFile,
        InvalidFileException,
        ParseError,
    ):
        return None


def _unsupported_secondary_bytes_reason(content: bytes) -> str:
    if content.startswith(b"PK\x03\x04"):
        # PK\x03\x04 is the shared ZIP magic for OOXML containers (pptx/xlsx/docx);
        # the byte prefix alone cannot distinguish them, so inspect the archive.
        return f"unsupported secondary document bytes format: {_ooxml_zip_kind(content)}"
    if content.startswith(b"%PDF-"):
        return "unsupported secondary document bytes format: pdf"
    return "unsupported secondary document bytes format: unknown"


def _ooxml_zip_kind(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        return "zip"
    if "ppt/presentation.xml" in names:
        return "pptx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "word/document.xml" in names:
        return "docx"
    return "zip"


def _fixture_performance_series() -> tuple[PerformanceSeries, PerformanceSeries, PerformanceSeries]:
    xlsx = PerformanceSeries(
        "monthly",
        (
            PerformancePoint(as_of=date(2025, 1, 31), value=0.021),
            PerformancePoint(as_of=date(2025, 2, 28), value=-0.012),
            PerformancePoint(as_of=date(2025, 3, 31), value=0.018),
            PerformancePoint(as_of=date(2025, 4, 30), value=0.006),
            PerformancePoint(as_of=date(2025, 5, 31), value=-0.004),
            PerformancePoint(as_of=date(2025, 6, 30), value=0.014),
        ),
    )
    deck = PerformanceSeries(
        "monthly",
        (
            PerformancePoint(as_of=date(2025, 1, 31), value=0.021),
            PerformancePoint(as_of=date(2025, 2, 28), value=-0.012),
            PerformancePoint(as_of=date(2025, 3, 31), value=0.025),
            PerformancePoint(as_of=date(2025, 4, 30), value=0.006),
            PerformancePoint(as_of=date(2025, 5, 31), value=-0.004),
            PerformancePoint(as_of=date(2025, 6, 30), value=0.014),
        ),
    )
    benchmark = PerformanceSeries(
        "monthly",
        (
            PerformancePoint(as_of=date(2025, 1, 31), value=0.010),
            PerformancePoint(as_of=date(2025, 2, 28), value=-0.008),
            PerformancePoint(as_of=date(2025, 3, 31), value=0.012),
            PerformancePoint(as_of=date(2025, 4, 30), value=0.003),
            PerformancePoint(as_of=date(2025, 5, 31), value=-0.006),
            PerformancePoint(as_of=date(2025, 6, 30), value=0.009),
        ),
    )
    return xlsx, deck, benchmark


#: Minimum-evidence thresholds for production (evidence-backed) scoring.
#: All five key fields must be present with provenance and confidence at or
#: above ``_MIN_FIELD_CONFIDENCE``; the workbook must yield at least
#: ``_MIN_WORKBOOK_OBSERVATIONS`` monthly observations with volatility,
#: drawdown and Sharpe available. Below any threshold, scoring stays
#: unavailable (``evidence_backed_score_components_unavailable``).
_MIN_FIELD_CONFIDENCE = 0.50
_MIN_WORKBOOK_OBSERVATIONS = 3
_KEY_FIELD_KEYS = (
    "strategy.asset_class",
    "terms.management_fee",
    "performance.net_return_1y",
    "operations.aum",
    "team.key_person_risk",
)


@dataclass(frozen=True)
class _DerivedScoring:
    asset_class: str
    components: tuple[ScoreComponent, ...]
    rationales: dict[str, str]
    provenances: dict[str, tuple[ComponentProvenance, ...]]


def _clamp01(value: float) -> float:
    return max(0.0, min(1.0, value))


def _team_experience_score(risk_value: str) -> float:
    """Score key-person risk text; adverse signals outrank reassuring ones.

    Matching is word-boundary based so substrings do not misfire (``high``
    in ``highly``, ``none`` in ``nonexistent``). A text that mentions both
    stability and a pending departure scores on the departure, not the
    reassurance.
    """

    risk_text = risk_value.lower()
    if re.search(r"\b(?:departure|pending|leaving|resigned)\b", risk_text):
        return 0.35
    if re.search(r"\bhigh\b", risk_text):
        return 0.40
    if re.search(r"\b(?:moderate|medium)\b", risk_text):
        return 0.60
    if (
        "no key person" in risk_text
        or "low risk" in risk_text
        or "stable team" in risk_text
        or re.search(r"\bnone\b", risk_text)
    ):
        return 0.90
    return 0.65


def _field_provenance(field: Any) -> str:
    return f"{field.key}@{field.source_doc_id}:page {field.source_page} (confidence {field.confidence:.2f})"


def _field_component_provenance(field: Any) -> ComponentProvenance:
    return ComponentProvenance(
        source_kind="document",
        reference=str(field.key),
        source_doc_id=str(field.source_doc_id) or None,
        source_page=int(field.source_page) if field.source_page is not None else None,
    )


def _workbook_component_provenance(workbook_ref: str) -> ComponentProvenance:
    file_name, separator, document_id = workbook_ref.rpartition("@")
    if not separator:
        return ComponentProvenance(source_kind="workbook", reference=workbook_ref)
    return ComponentProvenance(
        source_kind="workbook",
        reference=file_name,
        source_doc_id=document_id or None,
    )


def _parse_percent(value: str) -> float | None:
    match = re.search(r"(-?[0-9]+(?:\.[0-9]+)?)\s*%", value)
    return float(match.group(1)) / 100.0 if match else None


_AUM_AMOUNT = r"(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]+)?"
_AUM_WITH_UNIT = re.compile(
    rf"\$?\s*(?<![0-9.,])(?P<amount>{_AUM_AMOUNT})\s*"
    r"(?P<unit>billion|million|thousand|bn|mm|[kmb])\b",
    re.IGNORECASE,
)
_AUM_UNITLESS = re.compile(rf"\s*\$?\s*(?P<amount>{_AUM_AMOUNT})\s*\Z")


def _parse_aum(value: str) -> float | None:
    """Parse one unambiguous AUM amount without guessing from unrelated numbers."""

    candidates = list(_AUM_WITH_UNIT.finditer(value))
    if len(candidates) == 1:
        match = candidates[0]
        amount = float(match.group("amount").replace(",", ""))
        unit = match.group("unit").lower()
        multiplier = {
            "b": 1_000_000_000.0,
            "bn": 1_000_000_000.0,
            "billion": 1_000_000_000.0,
            "m": 1_000_000.0,
            "mm": 1_000_000.0,
            "million": 1_000_000.0,
            "k": 1_000.0,
            "thousand": 1_000.0,
        }[unit]
        return amount * multiplier
    if candidates:
        return None
    unitless = _AUM_UNITLESS.fullmatch(value)
    if unitless is None:
        return None
    return float(unitless.group("amount").replace(",", ""))


def _unavailable_explainability(reason: str) -> dict[str, object]:
    return {"status": "unavailable", "overall_score": None, "reason": reason}


def _evidence_backed_score_components(
    *,
    extraction: ExtractedDocumentResult,
    metrics: Any,
    workbook_ref: str,
) -> _DerivedScoring | None:
    """Map extracted fields and workbook metrics to the five score components.

    Returns ``None`` when evidence is below the minimum thresholds documented
    at ``_MIN_FIELD_CONFIDENCE`` / ``_MIN_WORKBOOK_OBSERVATIONS``.

    Component mapping (each in [0, 1], deterministic):
    - ``performance_consistency``: workbook volatility and max drawdown.
    - ``risk_adjusted_returns``: workbook Sharpe/Sortino mapped around 0.5,
      blended with the extracted 1Y net return.
    - ``operational_quality``: extracted AUM (log scale) and management fee.
    - ``transparency``: mean confidence of the five provenance-backed fields.
    - ``team_experience``: rule-based reading of the key-person-risk field.
    """

    fields = {field.key: field for field in extraction.fields}
    key_fields = []
    for key in _KEY_FIELD_KEYS:
        field = fields.get(key)
        if field is None or field.confidence < _MIN_FIELD_CONFIDENCE:
            return None
        if not field.source_doc_id or field.source_page is None or not field.value.strip():
            return None
        key_fields.append(field)
    if metrics is None or metrics.observation_count < _MIN_WORKBOOK_OBSERVATIONS:
        return None
    if (
        metrics.annualized_volatility is None
        or metrics.max_drawdown is None
        or metrics.sharpe_ratio is None
    ):
        return None
    try:
        asset_class = normalize_asset_class(fields["strategy.asset_class"].value)
    except ValueError:
        return None
    net_return = _parse_percent(fields["performance.net_return_1y"].value)
    fee = _parse_percent(fields["terms.management_fee"].value)
    aum = _parse_aum(fields["operations.aum"].value)
    if net_return is None or fee is None or aum is None or aum <= 0:
        return None

    by_key = {field.key: field for field in key_fields}
    performance_consistency = round(
        _clamp01(
            0.6 * (1.0 - metrics.annualized_volatility)
            + 0.4 * (1.0 - min(metrics.max_drawdown * 4.0, 1.0))
        ),
        6,
    )
    sharpe_like = metrics.sharpe_ratio
    if metrics.sortino_ratio is not None:
        sharpe_like = (metrics.sharpe_ratio + metrics.sortino_ratio) / 2.0
    risk_adjusted = round(
        _clamp01(0.8 * _clamp01(0.5 + sharpe_like / 4.0) + 0.2 * _clamp01(0.5 + net_return / 0.20)),
        6,
    )
    aum_score = _clamp01((math.log10(aum) - 6.0) / 4.0)
    operational_quality = round(_clamp01(0.6 * aum_score + 0.4 * _clamp01(1.0 - fee / 0.03)), 6)
    transparency = round(
        _clamp01(sum(field.confidence for field in key_fields) / len(key_fields)), 6
    )
    team_experience = _team_experience_score(fields["team.key_person_risk"].value)

    components = (
        ScoreComponent("performance_consistency", performance_consistency),
        ScoreComponent("risk_adjusted_returns", risk_adjusted),
        ScoreComponent("operational_quality", operational_quality),
        ScoreComponent("transparency", transparency),
        ScoreComponent("team_experience", team_experience),
    )
    rationales = {
        "performance_consistency": (
            f"Workbook volatility/drawdown from {workbook_ref} "
            f"(n={metrics.observation_count}); evidence {_field_provenance(by_key['performance.net_return_1y'])}"
        ),
        "risk_adjusted_returns": (
            f"Sharpe/Sortino from {workbook_ref} blended with extracted net return; "
            f"evidence {_field_provenance(by_key['performance.net_return_1y'])}"
        ),
        "operational_quality": (
            f"AUM and fee scale; evidence {_field_provenance(by_key['operations.aum'])}; "
            f"{_field_provenance(by_key['terms.management_fee'])}"
        ),
        "transparency": (
            "Mean confidence of provenance-backed key fields; evidence "
            + "; ".join(_field_provenance(field) for field in key_fields)
        ),
        "team_experience": (
            f"Key-person risk reading; evidence {_field_provenance(by_key['team.key_person_risk'])}"
        ),
    }
    workbook_provenance = _workbook_component_provenance(workbook_ref)
    provenances = {
        "performance_consistency": (
            workbook_provenance,
            _field_component_provenance(by_key["performance.net_return_1y"]),
        ),
        "risk_adjusted_returns": (
            workbook_provenance,
            _field_component_provenance(by_key["performance.net_return_1y"]),
        ),
        "operational_quality": (
            _field_component_provenance(by_key["operations.aum"]),
            _field_component_provenance(by_key["terms.management_fee"]),
        ),
        "transparency": tuple(_field_component_provenance(field) for field in key_fields),
        "team_experience": (_field_component_provenance(by_key["team.key_person_risk"]),),
    }
    return _DerivedScoring(
        asset_class=asset_class,
        components=components,
        rationales=rationales,
        provenances=provenances,
    )


def _workbook_provenance_ref(
    *,
    repository: CoreRepository,
    record: Any,
    file_entries: list[dict[str, Any]],
) -> str:
    file_name = _performance_track_record_file_name(file_entries) or "submitted_workbook"
    document_id = _document_id_for_file_name(
        repository=repository, document_ids=record.document_ids, file_name=file_name
    )
    if document_id is not None:
        return f"{file_name}@{document_id}"
    return file_name


def _score_components(benchmark_correlation: float | None) -> tuple[ScoreComponent, ...]:
    assert benchmark_correlation is not None
    correlation_component = max(0.0, min(1.0, benchmark_correlation))
    return (
        ScoreComponent("performance_consistency", 0.80),
        ScoreComponent("risk_adjusted_returns", 0.78),
        ScoreComponent("operational_quality", 0.69),
        ScoreComponent("transparency", 0.74),
        ScoreComponent("team_experience", round(correlation_component, 6)),
    )


def _explainability_inputs(
    asset_class: str,
    components: tuple[ScoreComponent, ...],
    rationales: dict[str, str] | None = None,
    provenances: dict[str, tuple[ComponentProvenance, ...]] | None = None,
) -> tuple[ScoreComponentInput, ...]:
    weights = get_weight_set(asset_class).weights
    rationale_by_component = rationales or {
        "performance_consistency": "Normalized monthly track record is complete.",
        "risk_adjusted_returns": "Prioritized metrics include Sharpe and correlation evidence.",
        "operational_quality": "AUM field is present but below auto-accept confidence.",
        "transparency": "Provenance pointers tie extracted fields to source pages.",
        "team_experience": "Key-person note is routed for analyst review.",
    }
    # Smoke/demo components have no submitted-evidence derivation; attribute
    # them to the analyst queue assignee instead of a document page.
    # Production callers pass ``provenances`` explicitly: a missing entry is a
    # derivation bug, not analyst evidence, so fail instead of misattributing.
    analyst_provenance = (ComponentProvenance(source_kind="analyst", reference="analyst_001"),)
    return tuple(
        ScoreComponentInput(
            component=component.name,
            weight=weights[component.name],
            score=component.value,
            rationale=rationale_by_component[component.name],
            provenance=(
                provenances[component.name] if provenances is not None else analyst_provenance
            ),
        )
        for component in components
    )


def _start_event(sink: InMemoryTraceSink, name: str) -> TraceEvent:
    matches = [event for event in sink.events if event.name == name and event.ended_at is None]
    assert matches, f"missing trace start event {name}"
    return matches[0]


def _derive_document_types(
    *, repository: CoreRepository, document_ids: tuple[str, ...]
) -> tuple[str, ...]:
    types: list[str] = []
    for document_id in document_ids:
        document = repository.get_document(document_id)
        if document is None:
            types.append("unknown")
            continue
        suffix = Path(document.file_name).suffix.lstrip(".").lower()
        types.append(suffix or "unknown")
    return tuple(sorted(set(types)))


def _document_source_names(
    *, repository: CoreRepository, document_ids: tuple[str, ...]
) -> tuple[str, ...]:
    names: list[str] = []
    for document_id in document_ids:
        document = repository.get_document(document_id)
        names.append(document.file_name if document is not None else document_id)
    return tuple(names)


def _pipeline_latency_ms(*, sink: InMemoryTraceSink, root_span_name: str) -> int | None:
    start_event = _start_event(sink, root_span_name)
    end_matches = [
        event
        for event in sink.events
        if event.name == root_span_name
        and event.span_id == start_event.span_id
        and event.ended_at is not None
    ]
    if not end_matches:
        return None
    started_at = _parse_iso_timestamp(start_event.started_at)
    ended_at = _parse_iso_timestamp(str(end_matches[0].ended_at))
    duration_ms = int((ended_at - started_at).total_seconds() * 1000)
    return max(duration_ms, 0)


def _parse_iso_timestamp(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def _assert_registered_package_state(
    *,
    service: IngestionService,
    package_id: str,
    expected_document_ids: tuple[str, ...],
) -> IngestRecord:
    record = service.get_record(package_id)
    assert record.status == "received"
    assert record.document_ids == expected_document_ids, "document identifiers must remain stable"
    assert service.get_events(package_id)[0].to_status == "received"
    return record
